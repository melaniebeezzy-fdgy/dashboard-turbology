#!/usr/bin/env python3
# Build D.stores roster for Mexico by reusing build_mx2.py store-assembly logic.
import json, re, unicodedata, os
LV = [1000, 2100, 2400, 2700, 3000]
def num(x):
    try: return float(x)
    except: return None
def to_m(x):
    v = num(x)
    return None if v is None else (v*1000 if v < 50 else v)
def sac(s): return ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn')
CITY_MAP = {'Ciudad de México': 'CDMX', 'Merida': 'Mérida'}
def ncity(c): return CITY_MAP.get(str(c), str(c))
def stripbrand(s): return re.sub(r'\s*-\s*Turbo\s*$', '', str(s or ''), flags=re.I).strip()

s2c = json.load(open('store2cocina_mx.json'))
coc2op = json.load(open('cocina2ops_mx.json')) if os.path.exists('cocina2ops_mx.json') else {}
def coc_of(sid): return s2c.get('MX'+sid) or s2c.get(sid) or s2c.get(sid+'.0') or '—'

RT = json.load(open('mx_rtwt_hist.json'))
POLY = json.load(open('mx_poly.json'))
MXW = RT['labels']
NW = len(MXW)
L = NW - 1

def _sidn(s):
    if s is None: return None
    s = re.sub(r'\.0+$', '', str(s).strip()); return re.sub(r'\D', '', s) or None
OLDPOLY = {}
if os.path.exists('FOODOLOGY.xlsx'):
    import openpyxl, warnings; warnings.filterwarnings('ignore')
    _wb = openpyxl.load_workbook('FOODOLOGY.xlsx', data_only=True); _ws = _wb['DETALLE']
    _h = {c: i for i, c in enumerate(next(_ws.iter_rows(min_row=1, max_row=1, values_only=True)))}
    for _r in _ws.iter_rows(min_row=2, values_only=True):
        _sid = _sidn(_r[_h['Store ID']]) if 'Store ID' in _h else None
        if not _sid: continue
        _cur = _r[_h['Current Size']] if 'Current Size' in _h else None
        _fin = _r[_h['FINAL SIZE']] if 'FINAL SIZE' in _h else None
        OLDPOLY[_sid] = (to_m(_cur), to_m(_fin))

stores = []
for sid, info in RT['stores'].items():
    coc = coc_of(sid)
    p = POLY.get(sid, {})
    city = ncity(p.get('city') or '—')
    op = coc2op.get(coc)
    rtarr = info['rt']
    pw = [None, None, None, None, to_m(p.get('a9')), to_m(p.get('a16')), to_m(p.get('a23')), to_m(p.get('a30'))]
    _op = OLDPOLY.get(sid)
    if _op:
        if _op[0] is not None: pw[1] = _op[0]
        if _op[1] is not None: pw[2] = _op[1]
    _last = None
    for _i in range(NW):
        if pw[_i] is not None: _last = pw[_i]
        elif _last is not None: pw[_i] = _last
    _first = next((v for v in pw if v is not None), None)
    pw = [v if v is not None else _first for v in pw]
    _pvals = [v for v in pw if v is not None]
    mxp = max(_pvals) if _pvals else None
    cc = next((pw[i] for i in range(NW-1, -1, -1) if pw[i] is not None), None)
    stores.append(dict(sid=sid, b=stripbrand(info.get('marca')), k=coc, c=city, op=op,
        rt=rtarr, fin=pw, cc=cc, mx=mxp))

if __name__ == '__main__' and os.environ.get('DUMP_ROSTER'):
    json.dump(stores, open('_dmx_roster.json', 'w'), ensure_ascii=False)
    print(len(stores), 'stores; rt-with-data', sum(1 for s in stores if any(v is not None for v in s['rt'])))
