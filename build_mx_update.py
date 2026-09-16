#!/usr/bin/env python3
# Actualiza México en el dashboard unificado (index_co.html) con un "Reporte Unificado" (.xlsx).
#
# USO:
#   python3 build_mx_update.py [ruta_reporte.xlsx] [ventas_semana.csv]
#     - Si no pasas el reporte, usa el más reciente "*Reporte_Unificado*.xlsx" en ../uploads.
#     - ventas_semana.csv (opcional): columnas store_id,q,gmv de la ÚLTIMA semana (la nueva).
#       Si no lo pasas, el script corre el histórico de ventas y deja la nueva semana en 0
#       (te recuerda el SQL para traerla del warehouse).
#
# Qué hace:
#   - Lee "Trazabilidad RTWT" (RTWT 8 semanas por tienda; las etiquetas salen de las fechas +6 = domingo).
#   - Lee "Polígonos" (FINAL SIZE, si no hay usa Current Size) como polígono actual por tienda.
#   - Toma la MXDATA actual del dashboard y alinea por ETIQUETA de semana: las semanas que ya existían
#     conservan su polígono/ventas; la(s) semana(s) nueva(s) usan FINAL SIZE y las ventas del CSV.
#   - Reinyecta const MXDATA y actualiza index.html (= copia de index_co.html).
import openpyxl, json, re, sys, os, glob, csv, datetime, unicodedata

HTML = 'index_co.html'
CITY_MAP = {'Ciudad de México': 'CDMX'}

def sac(s): return ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn')
def sidn(x):
    s = str(x if x is not None else '').strip()
    if s.endswith('.0'): s = s[:-2]
    s = s.upper().replace('MX', '')
    return re.sub(r'[^0-9]', '', s)
def num(x):
    try: return float(x)
    except Exception: return None
def stripbrand(s): return re.sub(r'\s*-?\s*Turbo.*$', '', str(s or ''), flags=re.I).strip()

# ---- args ----
args = [a for a in sys.argv[1:]]
report = next((a for a in args if a.lower().endswith('.xlsx')), None)
ventas_csv = next((a for a in args if a.lower().endswith('.csv')), None)
if not report:
    cands = sorted(glob.glob('../uploads/*Reporte_Unificado*.xlsx') + glob.glob('../uploads/*eporte*nificado*.xlsx'), key=os.path.getmtime)
    report = cands[-1] if cands else None
if not report or not os.path.exists(report):
    sys.exit('No encontré el reporte .xlsx. Pásalo como argumento.')
print('Reporte:', report)

wb = openpyxl.load_workbook(report, data_only=True)

# ---- Trazabilidad RTWT ----
tz = list(wb['Trazabilidad RTWT'].iter_rows(values_only=True))
hi = next(i for i, r in enumerate(tz) if r and str(r[0]).strip() == 'Store ID')
hdr = tz[hi]
def as_date(c):
    if isinstance(c, datetime.datetime): return c.date()
    if isinstance(c, datetime.date): return c
    if isinstance(c, str) and re.match(r'^\d{4}-\d{2}-\d{2}', c.strip()):
        return datetime.date.fromisoformat(c.strip()[:10])
    return None
wk_cols = [j for j, c in enumerate(hdr) if as_date(c) is not None]
def eow(c): return (as_date(c) + datetime.timedelta(days=6)).strftime('%b %-d')  # etiqueta = domingo
weeks = [eow(hdr[j]) for j in wk_cols]
NW = len(weeks)
rt_by = {}
for r in tz[hi+1:]:
    if not r or not r[0]: continue
    sd = sidn(r[0])
    rt_by[sd] = [num(r[j]) if j < len(r) else None for j in wk_cols]
print('Semanas del reporte:', weeks)

# ---- Polígonos (FINAL SIZE) ----
wp = list(wb['Polígonos'].iter_rows(values_only=True))
ph = wp[0]
def col(name): return ph.index(name) if name in ph else None
c_sid, c_fin, c_cur = col('Store ID'), col('FINAL SIZE'), col('Current Size')
c_brand, c_store, c_city, c_ctry = col('Brand Name'), col('Store Name'), col('City Name'), col('Country')
poly = {}
for r in wp[1:]:
    if c_ctry is not None and str(r[c_ctry]).strip().upper() != 'MX': continue
    sd = sidn(r[c_sid])
    if not sd: continue
    fin = num(r[c_fin]) if (c_fin is not None and r[c_fin] not in (None, '#REF!')) else None
    if fin is None and c_cur is not None: fin = num(r[c_cur])
    poly[sd] = dict(cc=fin, brand=r[c_brand] if c_brand is not None else None,
                    store=r[c_store] if c_store is not None else None,
                    city=r[c_city] if c_city is not None else None)

# ---- MXDATA actual ----
html = open(HTML, encoding='utf-8').read()
A = html.index('/*MXDATA_START*/const MXDATA=') + len('/*MXDATA_START*/const MXDATA=')
B = html.index('/*MXDATA_END*/')
MX = json.loads(html[A:B].rstrip().rstrip(';'))
oldD = MX['D']; old_weeks = oldD['weeks']
old_idx = {w: i for i, w in enumerate(old_weeks)}
old_by = {sidn(s['sid']): s for s in oldD['stores']}

# ---- ventas nueva semana (CSV opcional) ----
vent = {}
if not ventas_csv:
    for c in ['mx_ventas_lastweek.csv', '../uploads/mx_ventas_lastweek.csv']:
        if os.path.exists(c): ventas_csv = c; break
if ventas_csv and os.path.exists(ventas_csv):
    for row in csv.DictReader(open(ventas_csv)):
        sd = sidn(row.get('store_id') or row.get('Store ID') or '')
        if sd: vent[sd] = (int(float(row.get('q') or 0)), round(float(row.get('gmv') or 0)))
    print('Ventas nueva semana desde:', ventas_csv, '(', len(vent), 'tiendas )')
else:
    print('SIN ventas de la nueva semana. Corre este SQL en el warehouse, guárdalo como mx_ventas_lastweek.csv (store_id,q,gmv) y vuelve a correr:')
    lastmon = None
    # la última semana = último domingo; el order_week (lunes) = domingo-6
    print("  SELECT store_id, COUNT(*) q, ROUND(SUM(gmv)) gmv FROM fdgy_views.orders_consolidado")
    print("   WHERE country='MEX' AND provider_new_name='Rappi' AND brand ILIKE '%Turbo%' AND order_state='Finalized'")
    print("     AND order_week='<LUNES de la última semana, ej 2026-08-31>' GROUP BY store_id")

def cocina_guess(store):
    m = re.search(r'\(([^)]+)\)', str(store or ''))
    if m: return m.group(1).strip()
    t = sac(str(store or '')).split()
    return ' '.join(t[:2]) if t else 'Otra'

# ---- construir nuevas tiendas ----
sids = set(rt_by) | set(old_by)
newstores = []
LI = NW - 1
for sd in sids:
    old = old_by.get(sd)
    p = poly.get(sd, {})
    cc = p.get('cc') if p.get('cc') is not None else (old['cc'] if old else None)
    # meta
    if old:
        b, k, c, op = old['b'], old['k'], old['c'], old['op']
    else:
        b = stripbrand(p.get('brand') or p.get('store'))
        c = CITY_MAP.get(str(p.get('city')), p.get('city') or '—')
        k = cocina_guess(p.get('store')); op = None
    # rt
    if sd in rt_by:
        rt = rt_by[sd]
    elif old:
        rt = [old['rt'][old_idx[w]] if w in old_idx else None for w in weeks]
    else:
        rt = [None]*NW
    # fin/gw/ow alineados por etiqueta; semana nueva = cc / ventas
    fin, gw, ow = [], [], []
    for i, w in enumerate(weeks):
        if old and w in old_idx:
            j = old_idx[w]
            fin.append(old['fin'][j]); gw.append((old.get('gw') or [None]*len(old_weeks))[j]); ow.append((old.get('ow') or [None]*len(old_weeks))[j])
        else:  # semana nueva (o tienda nueva)
            fin.append(cc)
            q, g = vent.get(sd, (0, 0))
            gw.append(g); ow.append(q)
    last4 = [x for x in gw[-4:] if x]
    g = round(sum(last4)/len(last4)) if last4 else (old['g'] if old else None)
    newstores.append(dict(b=b, k=k, c=c, sid=sd, cc=cc, fin=fin, rt=rt, op=op, g=g, gw=gw, ow=ow))

MX['D'] = dict(weeks=weeks, n4=oldD.get('n4', 4), stores=newstores)
# MO se conserva (la vista mensual se calcula del histórico semanal por tienda)

payload = '/*MXDATA_START*/const MXDATA=' + json.dumps(MX, ensure_ascii=False) + ';/*MXDATA_END*/'
html = html[:html.index('/*MXDATA_START*/')] + payload + html[B+len('/*MXDATA_END*/'):]
open(HTML, 'w', encoding='utf-8').write(html)
open('index.html', 'w', encoding='utf-8').write(html)
nrt = sum(1 for s in newstores if s['rt'][LI] is not None)
ngw = sum(1 for s in newstores if (s['gw'][LI] or 0) > 0)
print('MX actualizado: %d tiendas | semanas %s | RTWT últ. sem: %d | ventas nueva sem: %d' % (len(newstores), weeks[0]+'..'+weeks[-1], nrt, ngw))
print('index.html regenerado. Haz: git add -A && git commit && git push')
